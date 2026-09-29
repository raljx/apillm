"""
Module d'extraction de texte depuis différents formats de documents.

Formats supportés :
- DOCX
- XLSX
- XLS
- PDF
- CSV

Le module est indépendant de Streamlit et FastAPI.
Il travaille uniquement avec :
- filename : nom du fichier
- content  : contenu binaire du fichier

Cela permet de l'utiliser aussi bien depuis :
- FastAPI
- des tests unitaires
- un worker asynchrone
- un script Python
"""

import io
import logging
import re
from typing import Final

import pandas as pd
from docx import Document
from pypdf import PdfReader


logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Extensions supportées
# ---------------------------------------------------------------------------

SUPPORTED_EXTENSIONS: Final[set[str]] = {
    ".docx",
    ".xlsx",
    ".xls",
    ".pdf",
    ".csv",
}


class FileExtractionError(Exception):
    """
    Exception levée lorsqu'un fichier ne peut pas être lu ou extrait.
    """


# ---------------------------------------------------------------------------
# Fonctions utilitaires
# ---------------------------------------------------------------------------

def get_file_extension(filename: str) -> str:
    """
    Retourne l'extension du fichier en minuscules.

    Exemple :
        "menu.PDF" -> ".pdf"

    Args:
        filename: Nom du fichier.

    Returns:
        Extension du fichier.

    Raises:
        FileExtractionError:
            Si le nom du fichier est absent ou invalide.
    """
    if not filename or not filename.strip():
        raise FileExtractionError(
            "Le nom du fichier est vide ou invalide."
        )

    filename = filename.strip()

    if "." not in filename:
        raise FileExtractionError(
            f"Impossible de déterminer l'extension du fichier : {filename}"
        )

    extension = "." + filename.rsplit(".", 1)[-1].lower()

    return extension


def validate_file(
    filename: str,
    content: bytes,
) -> str:
    """
    Vérifie que le fichier peut être traité.

    Args:
        filename: Nom du fichier.
        content: Contenu binaire du fichier.

    Returns:
        Extension normalisée du fichier.

    Raises:
        FileExtractionError:
            Si le fichier est vide ou non supporté.
    """
    if not content:
        raise FileExtractionError(
            f"Le fichier '{filename}' est vide."
        )

    extension = get_file_extension(filename)

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(
            sorted(SUPPORTED_EXTENSIONS)
        )

        raise FileExtractionError(
            f"Format non supporté pour '{filename}'. "
            f"Formats acceptés : {supported}"
        )

    return extension


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def extract_text_from_docx(content: bytes) -> str:
    """
    Extrait le texte d'un document Word DOCX.

    Le traitement récupère :
    - les paragraphes ;
    - les tableaux.

    Args:
        content: Contenu binaire du fichier DOCX.

    Returns:
        Texte extrait du document.

    Raises:
        FileExtractionError:
            En cas d'échec de lecture.
    """
    try:
        document = Document(
            io.BytesIO(content)
        )

        parts: list[str] = []

        # ---------------------------------------------------------------
        # Paragraphes
        # ---------------------------------------------------------------

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                parts.append(text)

        # ---------------------------------------------------------------
        # Tableaux
        # ---------------------------------------------------------------

        for table in document.tables:
            for row in table.rows:
                cells: list[str] = []

                for cell in row.cells:
                    cell_text = cell.text.strip()

                    if cell_text:
                        cells.append(cell_text)

                if cells:
                    # Séparateur explicite afin de conserver
                    # la structure logique des colonnes.
                    parts.append(
                        " | ".join(cells)
                    )

        extracted_text = "\n".join(parts)

        logger.info(
            "Extraction DOCX terminée : %s caractères.",
            len(extracted_text),
        )

        return extracted_text

    except Exception as exc:
        logger.exception(
            "Erreur lors de l'extraction d'un fichier DOCX."
        )

        raise FileExtractionError(
            f"Impossible de lire le document DOCX : {exc}"
        ) from exc


# ---------------------------------------------------------------------------
# EXCEL
# ---------------------------------------------------------------------------

def extract_text_from_excel(
    content: bytes,
    extension: str,
) -> str:
    """
    Extrait le contenu de tous les onglets d'un fichier Excel.

    Les cellules sont converties en texte avec conservation
    d'une structure de type tableau.

    Args:
        content:
            Contenu binaire du fichier.

        extension:
            Extension du fichier (.xlsx ou .xls).

    Returns:
        Texte extrait de tous les onglets.

    Raises:
        FileExtractionError:
            En cas d'échec de lecture.
    """
    try:
        if extension == ".xls":
            engine = "xlrd"
        else:
            engine = "openpyxl"

        excel_file = pd.ExcelFile(
            io.BytesIO(content),
            engine=engine,
        )

        parts: list[str] = []

        for sheet_name in excel_file.sheet_names:
            logger.info(
                "Extraction de l'onglet Excel : %s",
                sheet_name,
            )

            dataframe = pd.read_excel(
                excel_file,
                sheet_name=sheet_name,
                header=None,
                dtype=str,
            )

            # Remplacement des valeurs NaN par une chaîne vide.
            dataframe = dataframe.fillna("")

            parts.append(
                f"=== ONGLET : {sheet_name} ==="
            )

            for _, row in dataframe.iterrows():
                values: list[str] = []

                for value in row.tolist():
                    text = str(value).strip()

                    if text and text.lower() != "nan":
                        values.append(text)

                if values:
                    # Le séparateur permet au LLM de conserver
                    # la notion de colonnes.
                    parts.append(
                        " | ".join(values)
                    )

        extracted_text = "\n".join(parts)

        logger.info(
            "Extraction Excel terminée : %s caractères.",
            len(extracted_text),
        )

        return extracted_text

    except ImportError as exc:
        logger.exception(
            "Dépendance manquante pour la lecture Excel."
        )

        raise FileExtractionError(
            "Impossible de lire le fichier Excel. "
            "Vérifie que 'openpyxl' est installé pour les fichiers XLSX "
            "et 'xlrd' pour les fichiers XLS."
        ) from exc

    except Exception as exc:
        logger.exception(
            "Erreur lors de l'extraction du fichier Excel."
        )

        raise FileExtractionError(
            f"Impossible de lire le fichier Excel : {exc}"
        ) from exc


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def extract_text_from_pdf(content: bytes) -> str:
    """
    Extrait le texte d'un fichier PDF.

    Important :
    Cette fonction fonctionne pour les PDF contenant une couche texte.

    Pour un PDF scanné sous forme d'image, un OCR supplémentaire
    serait nécessaire.

    Args:
        content: Contenu binaire du PDF.

    Returns:
        Texte extrait du PDF.

    Raises:
        FileExtractionError:
            En cas d'échec de lecture.
    """
    try:
        reader = PdfReader(
            io.BytesIO(content)
        )

        parts: list[str] = []

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):
            page_text = page.extract_text() or ""

            page_text = page_text.strip()

            if page_text:
                parts.append(
                    f"=== PAGE {page_number} ==="
                )

                parts.append(page_text)

        extracted_text = "\n".join(parts)

        logger.info(
            "Extraction PDF terminée : %s pages, %s caractères.",
            len(reader.pages),
            len(extracted_text),
        )

        if not extracted_text.strip():
            logger.warning(
                "Le PDF ne contient aucun texte extractible."
            )

        return extracted_text

    except Exception as exc:
        logger.exception(
            "Erreur lors de l'extraction du fichier PDF."
        )

        raise FileExtractionError(
            f"Impossible de lire le fichier PDF : {exc}"
        ) from exc


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def extract_text_from_csv(content: bytes) -> str:
    """
    Extrait le contenu d'un fichier CSV.

    Plusieurs encodages et séparateurs sont essayés afin de gérer
    des fichiers provenant de sources différentes.

    Args:
        content: Contenu binaire du CSV.

    Returns:
        Texte structuré du CSV.

    Raises:
        FileExtractionError:
            Si aucun encodage ne permet de lire le fichier.
    """
    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "latin-1",
    ]

    last_error: Exception | None = None

    for encoding in encodings:
        try:
            dataframe = pd.read_csv(
                io.BytesIO(content),
                header=None,
                dtype=str,
                encoding=encoding,
                sep=None,
                engine="python",
            )

            dataframe = dataframe.fillna("")

            parts: list[str] = []

            for _, row in dataframe.iterrows():
                values: list[str] = []

                for value in row.tolist():
                    text = str(value).strip()

                    if text and text.lower() != "nan":
                        values.append(text)

                if values:
                    parts.append(
                        " | ".join(values)
                    )

            extracted_text = "\n".join(parts)

            logger.info(
                "Extraction CSV terminée avec encodage '%s' : "
                "%s caractères.",
                encoding,
                len(extracted_text),
            )

            return extracted_text

        except Exception as exc:
            last_error = exc

            logger.debug(
                "Échec de lecture CSV avec encodage '%s'.",
                encoding,
                exc_info=True,
            )

    raise FileExtractionError(
        "Impossible de lire le fichier CSV avec les encodages "
        f"testés. Dernière erreur : {last_error}"
    )


# ---------------------------------------------------------------------------
# NETTOYAGE DU TEXTE
# ---------------------------------------------------------------------------

def clean_extracted_text(text: str) -> str:
    """
    Nettoie le texte extrait avant son envoi au LLM.

    Objectifs :
    - supprimer les espaces superflus ;
    - normaliser les tabulations ;
    - réduire les lignes vides ;
    - supprimer les caractères invisibles courants ;
    - réduire la consommation inutile de tokens.

    Le nettoyage reste volontairement prudent afin de ne pas détruire
    la structure métier utile au modèle.

    Args:
        text: Texte brut extrait.

    Returns:
        Texte nettoyé.
    """
    if not text:
        return ""

    original_length = len(text)

    # ---------------------------------------------------------------
    # Normalisation des caractères invisibles
    # ---------------------------------------------------------------

    text = text.replace(
        "\u00A0",
        " ",
    )

    text = text.replace(
        "\u200B",
        "",
    )

    text = text.replace(
        "\uFEFF",
        "",
    )

    # ---------------------------------------------------------------
    # Normalisation des retours à la ligne
    # ---------------------------------------------------------------

    text = text.replace(
        "\r\n",
        "\n",
    )

    text = text.replace(
        "\r",
        "\n",
    )

    # ---------------------------------------------------------------
    # Normalisation des espaces et tabulations
    # ---------------------------------------------------------------

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    # Nettoyage des espaces autour des sauts de ligne.
    text = re.sub(
        r" *\n *",
        "\n",
        text,
    )

    # Maximum deux sauts de ligne consécutifs.
    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    # Nettoyage ligne par ligne.
    lines: list[str] = []

    for line in text.splitlines():
        line = line.strip()

        if line:
            lines.append(line)

    cleaned_text = "\n".join(lines).strip()

    logger.info(
        "Nettoyage du texte terminé : %s -> %s caractères.",
        original_length,
        len(cleaned_text),
    )

    return cleaned_text


# ---------------------------------------------------------------------------
# EXTRACTION PRINCIPALE
# ---------------------------------------------------------------------------

def extract_text(
    filename: str,
    content: bytes,
) -> str:
    """
    Point d'entrée principal du module.

    Cette fonction :
    1. valide le fichier ;
    2. détecte son extension ;
    3. sélectionne l'extracteur approprié ;
    4. extrait le texte ;
    5. nettoie le résultat ;
    6. vérifie que le document contient un contenu exploitable.

    Args:
        filename:
            Nom du fichier.

        content:
            Contenu binaire du fichier.

    Returns:
        Texte extrait et nettoyé.

    Raises:
        FileExtractionError:
            Si le fichier est invalide ou ne contient aucun texte exploitable.
    """
    extension = validate_file(
        filename=filename,
        content=content,
    )

    logger.info(
        "Début de l'extraction du fichier '%s' (%s).",
        filename,
        extension,
    )

    # ---------------------------------------------------------------
    # Sélection de l'extracteur
    # ---------------------------------------------------------------

    if extension == ".docx":
        raw_text = extract_text_from_docx(
            content
        )

    elif extension in {
        ".xlsx",
        ".xls",
    }:
        raw_text = extract_text_from_excel(
            content=content,
            extension=extension,
        )

    elif extension == ".pdf":
        raw_text = extract_text_from_pdf(
            content
        )

    elif extension == ".csv":
        raw_text = extract_text_from_csv(
            content
        )

    else:
        # Normalement impossible grâce à validate_file(),
        # mais protection supplémentaire.
        raise FileExtractionError(
            f"Aucun extracteur disponible pour '{extension}'."
        )

    # ---------------------------------------------------------------
    # Nettoyage
    # ---------------------------------------------------------------

    cleaned_text = clean_extracted_text(
        raw_text
    )

    # ---------------------------------------------------------------
    # Validation finale
    # ---------------------------------------------------------------

    if not cleaned_text:
        raise FileExtractionError(
            f"Aucun texte exploitable n'a été extrait du fichier "
            f"'{filename}'."
        )

    logger.info(
        "Extraction terminée pour '%s' : %s caractères utiles.",
        filename,
        len(cleaned_text),
    )

    return cleaned_text


# ---------------------------------------------------------------------------
# INFORMATIONS SUR LES FORMATS
# ---------------------------------------------------------------------------

def get_supported_extensions() -> list[str]:
    """
    Retourne la liste des extensions supportées.

    Cette fonction peut être utilisée par une route FastAPI pour exposer
    les capacités de l'API.

    Returns:
        Liste triée des extensions supportées.
    """
    return sorted(
        SUPPORTED_EXTENSIONS
    )